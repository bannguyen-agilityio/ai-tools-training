import csv
import io
import uuid
from datetime import date, datetime

import httpx
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncEngine

from tests.conftest import Login, MakeLogin
from tests.modules.balances.test_ledger_jobs import seed_year
from tests.modules.requests.test_requests import body, post

URL = "/api/v1/hr/reports/monthly"
NASTY = '=HYPERLINK("http://evil.test","x")'


def as_date(value: object) -> object:
    return value.date() if isinstance(value, datetime) else value


async def approve(owner: AsyncEngine, rid: str) -> None:
    async with owner.connect() as conn:
        await conn.execute(
            text("UPDATE leave_request SET status = 'approved' WHERE id = :r"),
            {"r": uuid.UUID(rid)},
        )


async def export(client: httpx.AsyncClient, who: Login, query: str) -> httpx.Response:
    return await client.get(f"{URL}{query}", headers=who.headers)


def csv_rows(r: httpx.Response) -> list[list[str]]:
    return list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))


async def test_requires_session(client: httpx.AsyncClient) -> None:
    assert (await client.get(f"{URL}?month=2026-06&format=csv")).status_code == 401


async def test_only_hr_admin_may_export(
    client: httpx.AsyncClient, make_login: MakeLogin, owner_engine: AsyncEngine
) -> None:
    deputy = await make_login()
    async with owner_engine.connect() as conn:
        await conn.execute(
            text("UPDATE employee SET is_hr_deputy = true WHERE id = :d"), {"d": deputy.employee_id}
        )
    for who in (await make_login(), await make_login("manager"), deputy):
        for fmt in ("csv", "xlsx"):
            assert (await export(client, who, f"?month=2026-06&format={fmt}")).status_code == 403
    async with owner_engine.connect() as conn:
        n = await conn.scalar(text("SELECT count(*) FROM audit_log WHERE action = 'report.export'"))
    assert n == 0  # refused calls are not exports


async def test_bad_parameters_are_422(client: httpx.AsyncClient, make_login: MakeLogin) -> None:
    hr = await make_login("hr")
    for q in ("", "?month=2026-13", "?month=0000-01", "?month=2026-06&format=pdf", "?month=x"):
        assert (await export(client, hr, q)).status_code == 422, q


async def test_csv_has_totals_balances_who_and_when(
    client: httpx.AsyncClient,
    staff: tuple[Login, Login],
    make_login: MakeLogin,
    owner_engine: AsyncEngine,
) -> None:
    emp, _ = staff
    hr = await make_login("hr")
    await seed_year(owner_engine, emp.employee_id, "entitlement", "7", type="sick")
    annual = (await post(client, emp, body("2026-06-08", "2026-06-10"))).json()["id"]  # 3 days
    sick = (
        await post(
            client, emp, body("2026-06-15", leave_type="sick", start_part="am", end_part="am")
        )
    ).json()["id"]
    await approve(owner_engine, annual)
    await approve(owner_engine, sick)
    pending = (await post(client, emp, body("2026-06-22"))).json()["id"]  # not leave taken
    gone = (await post(client, emp, body("2026-06-23"))).json()["id"]
    await post(client, emp, {"version": 1}, f"/api/v1/requests/{gone}/cancel")
    assert pending

    r = await export(client, hr, "?month=2026-06&format=csv")
    assert r.status_code == 200
    assert r.headers["content-disposition"] == 'attachment; filename="leave-report-2026-06.csv"'
    assert r.headers["cache-control"] == "no-store" and r.headers["x-confidential"] == "HR only"
    rows = csv_rows(r)
    assert rows[0][0].startswith("CONFIDENTIAL")
    flat = [tuple(x) for x in rows]
    name = "Test"
    assert (name, "annual", "3.0") in flat
    assert (name, "sick", "0.5") in flat  # HR sees the real type
    assert not any(row[:1] == (name,) and row[1:2] == ("unpaid",) for row in flat)
    assert (name, "annual", "2026-06-08", "2026-06-10", "3.0") in flat
    assert (name, "sick", "2026-06-15", "2026-06-15", "0.5") in flat
    assert (name, "annual", "2026-06-22", "2026-06-22", "1.0") not in flat  # Pending
    # End of June: 12 granted less 3 taken; the Pending request is not counted; Sick 7 less half.
    assert (name, "9.0", "0", "6.5") in flat


async def test_xlsx_is_marked_confidential_and_matches_the_csv(
    client: httpx.AsyncClient,
    staff: tuple[Login, Login],
    make_login: MakeLogin,
    owner_engine: AsyncEngine,
) -> None:
    emp, _ = staff
    hr = await make_login("hr")
    rid = (await post(client, emp, body("2026-06-08"))).json()["id"]
    await approve(owner_engine, rid)
    r = await export(client, hr, "?month=2026-06")  # xlsx is the default
    assert r.status_code == 200
    assert r.headers["content-disposition"] == 'attachment; filename="leave-report-2026-06.xlsx"'
    book = load_workbook(io.BytesIO(r.content))
    assert book.sheetnames == ["Days taken", "Balances", "Who was off"]
    assert "CONFIDENTIAL" in (book.properties.title or "")
    assert book.properties.subject == "CONFIDENTIAL - HR only"
    for sheet in book.worksheets:
        assert str(sheet["A1"].value).startswith("CONFIDENTIAL")
    days = [tuple(c.value for c in row) for row in book["Days taken"].iter_rows(min_row=4)]
    assert ("Test", "annual", 1.0) in days
    off = [
        (a, b, as_date(c), as_date(d), e)
        for a, b, c, d, e in (
            tuple(c.value for c in row) for row in book["Who was off"].iter_rows(min_row=4)
        )
    ]
    assert ("Test", "annual", date(2026, 6, 8), date(2026, 6, 8), 1.0) in off


async def test_formula_text_is_neutralised_in_both_formats(
    client: httpx.AsyncClient,
    staff: tuple[Login, Login],
    make_login: MakeLogin,
    owner_engine: AsyncEngine,
) -> None:
    emp, _ = staff
    hr = await make_login("hr")
    async with owner_engine.connect() as conn:
        await conn.execute(
            text("UPDATE employee SET name = :n WHERE id = :e"), {"n": NASTY, "e": emp.employee_id}
        )
    rid = (await post(client, emp, body("2026-06-08"))).json()["id"]
    await approve(owner_engine, rid)

    got = csv_rows(await export(client, hr, "?month=2026-06&format=csv"))
    cells = [c for row in got for c in row]
    assert "'" + NASTY in cells and NASTY not in cells
    assert not any(c.startswith(("=", "+", "-", "@", "\t")) for c in cells if c)

    book = load_workbook(
        io.BytesIO((await export(client, hr, "?month=2026-06&format=xlsx")).content)
    )
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                assert cell.data_type != "f", (sheet.title, cell.coordinate)
                if isinstance(cell.value, str):
                    assert not cell.value.startswith(("=", "+", "-", "@", "\t"))
    names = [c.value for c in book["Days taken"]["A"]]
    assert "'" + NASTY in names


async def test_each_export_is_audited_without_names_or_types(
    client: httpx.AsyncClient,
    staff: tuple[Login, Login],
    make_login: MakeLogin,
    owner_engine: AsyncEngine,
) -> None:
    emp, _ = staff
    hr = await make_login("hr")
    await seed_year(owner_engine, emp.employee_id, "entitlement", "7", type="sick")
    rid = (await post(client, emp, body("2026-06-15", leave_type="sick"))).json()["id"]
    await approve(owner_engine, rid)
    await export(client, hr, "?month=2026-06&format=csv")
    await export(client, hr, "?month=2026-06&format=xlsx")
    async with owner_engine.connect() as conn:
        rows = (
            await conn.execute(
                text(
                    "SELECT actor_id, target_id, detail::text FROM audit_log "
                    "WHERE action = 'report.export' ORDER BY at"
                )
            )
        ).all()
    assert len(rows) == 2
    for actor, target, detail in rows:
        assert actor == hr.employee_id and target == "monthly-2026-06"
        assert "2026-06" in detail and "sick" not in detail.lower() and "Test" not in detail
    assert '"csv"' in rows[0][2] and '"xlsx"' in rows[1][2]


async def test_balances_ignore_leave_after_the_month_and_equal_names_stay_apart(
    client: httpx.AsyncClient,
    staff: tuple[Login, Login],
    make_login: MakeLogin,
    owner_engine: AsyncEngine,
) -> None:
    emp, mgr = staff
    twin = await make_login()
    hr = await make_login("hr")
    async with owner_engine.connect() as conn:
        await conn.execute(
            text("UPDATE employee SET manager_id = :m WHERE id = :t"),
            {"m": mgr.employee_id, "t": twin.employee_id},
        )
    await seed_year(owner_engine, twin.employee_id, "entitlement", "12")
    for who, start in ((emp, "2026-06-08"), (twin, "2026-06-09")):
        await approve(owner_engine, (await post(client, who, body(start))).json()["id"])
    later = (await post(client, emp, body("2026-12-07", "2026-12-09"))).json()["id"]  # after June
    await approve(owner_engine, later)

    flat = [tuple(x) for x in csv_rows(await export(client, hr, "?month=2026-06&format=csv"))]
    assert flat.count(("Test", "annual", "1.0")) == 2  # two people, not one row of 2.0
    annual_balances = [
        r for r in flat if len(r) == 4 and r[0] == "Test" and r[1] != "Annual carried"
    ]
    assert ("Test", "11.0", "0", "0") in annual_balances  # December leave does not reduce June
